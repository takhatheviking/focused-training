"""Event reports: participants and orders as tables, written to .xlsx here and to Google
Sheets by app.services.gsheets. Both outputs share the same rows, colours and layout."""

import io
from dataclasses import dataclass, field
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Event, Order, Payment, PaymentStatus, Ticket
from app.formatting import fmt_dt, fmt_money
from app.i18n import t
from app.services.events import event_stats

ACCENT = "4F46E5"
BAND = "F3F4F6"
GRID = Side(style="thin", color="E5E7EB")
DATE_FORMAT = "DD.MM.YYYY HH:MM"
MONEY_FORMAT = "#,##0"
HEADER_ROW = 5  # 1-based; rows 1-3 hold the title block

# (fill, font colour) per status; anything else stays plain.
TICKET_STYLE = {
    "used": ("DCFCE7", "166534"),
    "cancelled": ("FEE2E2", "991B1B"),
}
ORDER_STYLE = {
    "paid": ("DCFCE7", "166534"),
    "under_review": ("FEF3C7", "92400E"),
    "pending": ("DBEAFE", "1E40AF"),
    "refund_required": ("FEE2E2", "991B1B"),
    "refunded": ("F3F4F6", "4B5563"),
    "expired": ("F3F4F6", "6B7280"),
    "cancelled": ("F3F4F6", "6B7280"),
}


@dataclass
class Table:
    # Title, subtitle and a summary line; writers append the "exported at" stamp to the last one.
    title_lines: list[str]
    headers: list[str]
    widths: list[int]  # in characters
    rows: list[list]  # str | int | float | datetime | None
    status_col: int  # 1-based
    status_keys: list[str]  # raw status per row, for colouring
    styles: dict[str, tuple[str, str]]
    status_labels: dict[str, str]  # raw status -> shown text (for conditional formatting)
    kinds: dict[int, str] = field(default_factory=dict)  # 1-based col -> "date" | "money"


def _local(dt: datetime | None, tz: ZoneInfo) -> datetime | None:
    # Spreadsheets have no time zones: store the Tashkent wall-clock time.
    return dt.astimezone(tz).replace(tzinfo=None) if dt else None


def _sums(tiyin: int) -> int | float:
    return tiyin // 100 if tiyin % 100 == 0 else tiyin / 100


def _event_header(event: Event, lang: str, tz: ZoneInfo) -> tuple[str, str]:
    place = event.loc("venue", lang) or ""
    return event.loc("title", lang), " · ".join(p for p in (fmt_dt(event.starts_at, tz, lang), place) if p)


def _generated(lang: str, tz: ZoneInfo) -> str:
    return t(lang, "xls_generated", at=fmt_dt(datetime.now(tz), tz, lang, with_weekday=False))


def stamped_title_lines(table: Table, lang: str, tz: ZoneInfo) -> list[str]:
    *head, last = table.title_lines
    return [*head, " · ".join(p for p in (last, _generated(lang, tz)) if p)]


async def participants_table(session: AsyncSession, event: Event, lang: str, tz: ZoneInfo) -> Table:
    stats = await event_stats(session, event)
    title, subtitle = _event_header(event, lang, tz)
    summary = t(lang, "xls_summary", sold=stats.sold, checked=stats.checked_in, revenue=fmt_money(stats.revenue_tiyin, lang))
    tickets = list(
        (await session.execute(select(Ticket).where(Ticket.event_id == event.id).order_by(Ticket.ticket_type_id, Ticket.id))).scalars()
    )
    paid_at = dict((await session.execute(select(Order.id, Order.paid_at).where(Order.event_id == event.id))).all())
    return Table(
        title_lines=[title, subtitle, summary],
        headers=t(lang, "xls_participants_headers").split("|"),
        widths=[6, 13, 28, 17, 18, 16, 15, 17, 10, 17],
        rows=[
            [
                i,
                tk.code,
                tk.holder_name,
                tk.order.buyer_phone or "",
                f"@{tk.order.buyer_username}" if tk.order.buyer_username else "",
                tk.ticket_type.loc("name", lang),
                t(lang, f"xls_ticket_{tk.status}"),
                _local(tk.checked_in_at, tz),
                tk.order_id,
                _local(paid_at.get(tk.order_id), tz),
            ]
            for i, tk in enumerate(tickets, start=1)
        ],
        status_col=7,
        status_keys=[tk.status for tk in tickets],
        styles=TICKET_STYLE,
        status_labels={s: t(lang, f"xls_ticket_{s}") for s in ("valid", "used", "cancelled")},
        kinds={8: "date", 10: "date"},
    )


async def orders_table(session: AsyncSession, lang: str, tz: ZoneInfo, event: Event | None = None, limit: int = 5000) -> Table:
    """Orders of one event, or of all events (then with an extra "Event" column)."""
    stmt = select(Order)
    stmt = stmt.where(Order.event_id == event.id).order_by(Order.id) if event else stmt.order_by(Order.id.desc()).limit(limit)
    orders = list((await session.execute(stmt)).scalars())
    reviewers = dict(
        (
            await session.execute(
                select(Payment.order_id, Payment.reviewed_by_name).where(
                    Payment.order_id.in_([o.id for o in orders]), Payment.status == PaymentStatus.APPROVED
                )
            )
        ).all()
    ) if orders else {}

    headers = t(lang, "xls_orders_headers").split("|")
    widths = [10, 17, 26, 17, 18, 16, 8, 14, 10, 18, 17, 17, 24]
    rows = [
        [
            o.id,
            t(lang, f"order_status_{o.status}"),
            o.buyer_name,
            o.buyer_phone or "",
            f"@{o.buyer_username}" if o.buyer_username else "",
            o.ticket_type.loc("name", lang),
            o.quantity,
            _sums(o.amount_tiyin),
            f"•••• {o.card_last4}" if o.card_last4 else "",
            t(lang, f"xls_method_{o.payment_method}"),
            _local(o.created_at, tz),
            _local(o.paid_at, tz),
            reviewers.get(o.id) or (o.issued_by.full_name if o.issued_by else ""),
        ]
        for o in orders
    ]
    kinds = {8: "money", 11: "date", 12: "date"}
    status_col = 2
    if event is None:
        # One sheet for all events: the event goes into the second column.
        headers.insert(1, t(lang, "xls_col_event"))
        widths.insert(1, 34)
        for o, row in zip(orders, rows):
            row.insert(1, f"{o.event.starts_at.astimezone(tz):%d.%m} · {o.event.loc('title', lang)}")
        kinds = {col + 1: kind for col, kind in kinds.items()}
        status_col = 3
        title_lines = [t(lang, "gs_tab_orders"), "", ""]
    else:
        title_lines = [*_event_header(event, lang, tz), ""]
    return Table(
        title_lines=title_lines,
        headers=headers,
        widths=widths,
        rows=rows,
        status_col=status_col,
        status_keys=[o.status for o in orders],
        styles=ORDER_STYLE,
        status_labels={s: t(lang, f"order_status_{s}") for s in ORDER_STYLE},
        kinds=kinds,
    )


# --- .xlsx ------------------------------------------------------------------------------

def _set(ws: Worksheet, row: int, col: int, value):
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # never a formula, even if a name starts with "="
    return cell


def write_sheet(ws: Worksheet, table: Table, lang: str, tz: ZoneInfo) -> None:
    """Title block, coloured header, banded rows, status colours, filter, frozen header."""
    ws.sheet_view.showGridLines = False
    for i, line in enumerate(stamped_title_lines(table, lang, tz), start=1):
        cell = _set(ws, i, 1, line)
        cell.font = Font(size=14, bold=True, color="111827") if i == 1 else Font(size=10, color="6B7280")
    for col, (name, width) in enumerate(zip(table.headers, table.widths), start=1):
        cell = _set(ws, HEADER_ROW, col, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[HEADER_ROW].height = 24

    for r, (values, status_key) in enumerate(zip(table.rows, table.status_keys), start=HEADER_ROW + 1):
        band = PatternFill("solid", fgColor=BAND) if (r - HEADER_ROW) % 2 == 0 else None
        for col, value in enumerate(values, start=1):
            cell = _set(ws, r, col, value)
            cell.border = Border(bottom=GRID)
            cell.alignment = Alignment(vertical="center")
            if band:
                cell.fill = band
            kind = table.kinds.get(col)
            if kind == "date" and value is not None:
                cell.number_format = DATE_FORMAT
            elif kind == "money" and value is not None:
                cell.number_format = MONEY_FORMAT
        if status_key in table.styles:
            fill, color = table.styles[status_key]
            cell = ws.cell(row=r, column=table.status_col)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(bold=True, color=color)

    last_row = HEADER_ROW + max(len(table.rows), 1)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(table.headers))}{last_row}"
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)


async def build_event_report(session: AsyncSession, event: Event, lang: str, tz: ZoneInfo) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = t(lang, "xls_sheet_participants")[:31]
    write_sheet(ws, await participants_table(session, event, lang, tz), lang, tz)
    write_sheet(wb.create_sheet(t(lang, "xls_sheet_orders")[:31]), await orders_table(session, lang, tz, event), lang, tz)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
